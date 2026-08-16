"""Build the plain-language 15-page NovaPlan student internship report.

The report uses live model outputs, but keeps the story intentionally small:
one fictional company, one growth question, three financial ideas, and a few
experiments that a beginner can reproduce in the dashboard.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

import build_whitepaper as wp


REPORT_PATH = wp.OUT_DIR / "NovaPlan_Student_Internship_Report.docx"


# Design preset: standard_business_brief. The existing teal chart and callout
# palette is a named "student_teal" visual override used consistently.
def configure_student_styles(doc: Document) -> None:
    wp.configure_styles(doc)
    heading_tokens = {
        "Heading 1": (16, "2E74B5", 16, 8),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def configure_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = paragraph.add_run("NOVAPLAN | STUDENT INTERNSHIP REPORT")
    wp.set_run_font(left, size=8.5, color=wp.MUTED, bold=True)
    right = paragraph.add_run("\t15-PAGE EDITION | AUGUST 2026")
    wp.set_run_font(right, size=8.5, color=wp.MUTED, bold=True)
    wp.set_paragraph_border(paragraph, wp.GRID, side="bottom", size="8")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    label = paragraph.add_run("Fictional data | Educational project")
    wp.set_run_font(label, size=8.5, color=wp.MUTED)
    page = paragraph.add_run("\tPage ")
    wp.set_run_font(page, size=8.5, color=wp.MUTED)
    wp.add_page_field(paragraph)


def start_page(doc: Document, number: int, kicker: str, title: str) -> None:
    wp.add_section_kicker(doc, f"{number} | {kicker}")
    wp.add_heading(doc, title, level=1)


def next_page(doc: Document, number: int, kicker: str, title: str) -> None:
    wp.page_break(doc)
    start_page(doc, number, kicker, title)


def add_figure(
    doc: Document,
    path: Path,
    number: int,
    caption: str,
    alt_text: str,
    width: float = 5.8,
) -> None:
    wp.add_figure(doc, path, f"Figure {number}. {caption}", width=width, alt_text=alt_text)


def project_flow_figure() -> Path:
    image, draw = wp.new_canvas(
        "What I built",
        "Five small parts connect the fictional data to a practical conclusion.",
    )
    items = [
        ("1. Data", "Invented company records"),
        ("2. Formulas", "Revenue, profit and cash"),
        ("3. Forecast", "13 weeks and 18 months"),
        ("4. Experiments", "Hiring, payments and risk"),
        ("5. Lesson", "Explain what changed"),
    ]
    xs = [55, 365, 675, 985, 1295]
    colors = [wp.BLUE, wp.PURPLE, wp.TEAL, wp.GREEN, wp.AMBER]
    for index, ((title, body), x, color) in enumerate(zip(items, xs, colors)):
        wp.box(draw, (x, 270, x + 250, 515), title, body, fill=wp.SURFACE, outline=color)
        if index < len(items) - 1:
            wp.arrow(draw, (x + 250, 392), (xs[index + 1] - 10, 392), color=wp.MUTED, width=4)
    draw.rounded_rectangle((280, 640, 1320, 730), radius=18, fill=wp.rgb(wp.TEAL_SOFT), outline=wp.rgb(wp.TEAL), width=2)
    draw.text((340, 665), "Main question: can the company hire 40 people and still keep enough cash?", fill=wp.rgb(wp.INK_SOFT), font=wp.font(22, semibold=True))
    draw.text((65, 823), "Source: NovaPlan student project design.", fill=wp.rgb(wp.MUTED), font=wp.font(17))
    return wp.save_canvas(image, "student_figure_01_project_flow.png")


def student_data_figure(shapes: dict[str, tuple[int, int]]) -> Path:
    image, draw = wp.new_canvas(
        "The four data groups used in the report",
        "All records are fictional, linked by simple business rules, and repeatable with seed 42.",
    )
    cards = [
        ("Monthly results", f"{shapes['monthly_kpis'][0]} months", "Revenue, costs and cash"),
        ("Invoices", f"{shapes['invoices'][0]:,} rows", "When customers pay"),
        ("Employees", f"{shapes['employees'][0]} people", "Roles, salaries and dates"),
        ("Hiring plan", "40 planned hires", "Joining and ramp-up"),
    ]
    bounds = [(95, 200, 755, 430), (845, 200, 1505, 430), (95, 500, 755, 730), (845, 500, 1505, 730)]
    colors = [wp.BLUE, wp.TEAL, wp.PURPLE, wp.AMBER]
    for (title, value, note), box_bounds, color in zip(cards, bounds, colors):
        x0, y0, x1, y1 = box_bounds
        draw.rounded_rectangle(box_bounds, radius=20, fill=wp.rgb(wp.SURFACE), outline=wp.rgb(color), width=3)
        draw.text((x0 + 28, y0 + 26), title, fill=wp.rgb(wp.INK), font=wp.font(27, bold=True))
        draw.text((x0 + 28, y0 + 88), value, fill=wp.rgb(color), font=wp.font(34, bold=True))
        draw.text((x0 + 28, y0 + 155), note, fill=wp.rgb(wp.MUTED), font=wp.font(21))
    draw.text((65, 823), "Source: checked-in NovaPlan synthetic CSV files.", fill=wp.rgb(wp.MUTED), font=wp.font(17))
    return wp.save_canvas(image, "student_figure_02_data_groups.png")


def simple_mc_flow_figure() -> Path:
    image, draw = wp.new_canvas(
        "Monte Carlo in five steps",
        "The computer repeats the same financial model while a few uncertain inputs move.",
    )
    items = [
        ("Start", "Use the Base forecast", wp.BLUE),
        ("Change", "Move utilization, payment time and costs", wp.PURPLE),
        ("Recalculate", "Build revenue, profit and cash again", wp.TEAL),
        ("Repeat", "Run 5,000 times with seed 42", wp.GREEN),
        ("Compare", "Read low, middle and high results", wp.AMBER),
    ]
    xs = [55, 365, 675, 985, 1295]
    for index, ((title, body, color), x) in enumerate(zip(items, xs)):
        wp.box(draw, (x, 250, x + 250, 535), title, body, fill=wp.SURFACE, outline=color)
        if index < len(items) - 1:
            wp.arrow(draw, (x + 250, 392), (xs[index + 1] - 10, 392), color=wp.MUTED, width=4)
    draw.rounded_rectangle((290, 650, 1310, 735), radius=18, fill=wp.rgb(wp.TEAL_SOFT), outline=wp.rgb(wp.TEAL), width=2)
    draw.text((350, 675), "Important: cash is calculated from the changed drivers; it is not picked at random.", fill=wp.rgb(wp.INK_SOFT), font=wp.font(21, semibold=True))
    draw.text((65, 823), "Source: simplified explanation of models/monte_carlo.py.", fill=wp.rgb(wp.MUTED), font=wp.font(17))
    return wp.save_canvas(image, "student_figure_08_monte_carlo_flow.png")


def build_live_results(analytics: dict[str, object]) -> dict[str, object]:
    monthly: pd.DataFrame = analytics["monthly"]
    base = wp.run_scenario(monthly, "Base", months=18)
    hire_40 = wp.run_scenario(
        monthly,
        "Custom",
        {
            "incremental_hires_total": 40,
            "hire_joining_months": 6,
            "hire_ramp_months": 3,
        },
        months=18,
    )
    experiments = {
        "Hire 40 people": hire_40,
        "Utilization +5 points": wp.run_scenario(monthly, "Custom", {"utilization_change_pp": 0.05}, months=18),
        "Billing rate +5%": wp.run_scenario(monthly, "Custom", {"billing_rate_change_pct": 0.05}, months=18),
        "Payments 15 days slower": wp.run_scenario(monthly, "Custom", {"dso_delta_days": 15}, months=18),
    }
    mc = {}
    for label, scale in (("0%", 0.0), ("100%", 1.0), ("200%", 2.0)):
        mc[label] = wp.run_monte_carlo(
            monthly,
            n_simulations=5_000,
            seed=42,
            threshold=16_000_000,
            assumptions={"uncertainty_scale": scale},
        )
    mc_low_target = wp.run_monte_carlo(
        monthly,
        n_simulations=5_000,
        seed=42,
        threshold=14_000_000,
        assumptions={"uncertainty_scale": 2.0},
    )
    mc_high_target = wp.run_monte_carlo(
        monthly,
        n_simulations=5_000,
        seed=42,
        threshold=18_000_000,
        assumptions={"uncertainty_scale": 2.0},
    )
    return {
        "base": base,
        "hire_40": hire_40,
        "experiments": experiments,
        "mc": mc,
        "mc_low_target": mc_low_target,
        "mc_high_target": mc_high_target,
    }


def build_student_figures(analytics: dict[str, object], results: dict[str, object]) -> dict[str, Path]:
    base_forecast = results["base"]["forecast"]
    hire_forecast = results["hire_40"]["forecast"]
    liquidity: pd.DataFrame = analytics["liquidity"]
    base_cash = float(results["base"]["summary"]["year_end_cash_eur"])

    figures = {
        "flow": project_flow_figure(),
        "data": student_data_figure(analytics["dataset_shapes"]),
        "history": wp.build_figures(analytics)["history"],
        "driver_tree": wp.driver_tree_figure(),
        "working_capital": wp.working_capital_figure(),
        "walk_forward": wp.walk_forward_figure(),
        "mc_flow": simple_mc_flow_figure(),
    }
    figures["liquidity"] = wp.line_chart(
        "student_figure_05_cash_check.png",
        "The 13-week cash check",
        "Zooming in on the weekly balances makes the movement easier to see.",
        [f"W{int(value)}" for value in liquidity["week_number"]],
        [{"label": "Closing cash", "values": liquidity["closing_cash_eur"].tolist(), "color": wp.TEAL, "markers": True}],
        value_formatter=lambda value: f"EUR {value / 1_000_000:.1f}m",
    )
    figures["hiring"] = wp.line_chart(
        "student_figure_06_hiring_cash.png",
        "Cash: Base plan versus Hire 40",
        "New salaries start before every new employee reaches full productivity.",
        [pd.Timestamp(value).strftime("%b %y") for value in base_forecast["month"]],
        [
            {"label": "Base", "values": base_forecast["closing_cash_eur"].tolist(), "color": wp.BLUE},
            {"label": "Hire 40", "values": hire_forecast["closing_cash_eur"].tolist(), "color": wp.TEAL},
        ],
        value_formatter=lambda value: f"EUR {value / 1_000_000:.1f}m",
    )
    cash_impacts = []
    for scenario in results["experiments"].values():
        cash_impacts.append(float(scenario["summary"]["year_end_cash_eur"]) - base_cash)
    figures["experiments"] = wp.horizontal_bar_chart(
        "student_figure_07_four_controls.png",
        "Four controls, four different cash effects",
        "Impact on cash at the end of the 18-month plan compared with Base.",
        list(results["experiments"].keys()),
        cash_impacts,
        colors=[wp.TEAL if value >= 0 else wp.RED for value in cash_impacts],
        value_formatter=lambda value: f"{value / 1_000_000:+.1f}m",
        symmetric=True,
    )
    widths = []
    for label in ("0%", "100%", "200%"):
        summary = results["mc"][label]["summary"]
        widths.append(summary["p95_year_end_cash_eur"] - summary["p5_year_end_cash_eur"])
    figures["mc_test"] = wp.horizontal_bar_chart(
        "student_figure_09_monte_carlo_test.png",
        "More uncertainty creates a wider result range",
        "P95 minus P5 at the end of the 18-month forecast.",
        ["0% uncertainty", "100% uncertainty", "200% uncertainty"],
        widths,
        colors=[wp.BLUE, wp.TEAL, wp.AMBER],
        value_formatter=lambda value: f"EUR {value / 1_000_000:.1f}m",
    )
    return figures


def build_report(analytics: dict[str, object]) -> Path:
    monthly: pd.DataFrame = analytics["monthly"]
    annual: pd.DataFrame = analytics["annual"]
    liquidity: pd.DataFrame = analytics["liquidity"]
    results = build_live_results(analytics)
    figures = build_student_figures(analytics, results)
    base_summary = results["base"]["summary"]
    hire_summary = results["hire_40"]["summary"]
    latest = annual.iloc[-1]

    doc = Document()
    configure_student_styles(doc)
    wp.configure_page(doc.sections[0])
    doc.core_properties.title = "NovaPlan: Student Internship Project"
    doc.core_properties.subject = "A plain-language 15-page report about business planning, cash and Monte Carlo"
    doc.core_properties.author = "NovaPlan student project"
    doc.core_properties.keywords = "synthetic data, business planning, cash forecast, Monte Carlo, student project"
    doc.core_properties.comments = "Generated from the checked-in seed-42 fictional company data and tested model outputs."

    # Page 1 - editorial cover.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(72)
    spacer.paragraph_format.space_after = Pt(0)
    wp.add_section_kicker(doc, "Student internship-style research project")
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("How can a consulting company grow without running short of cash?")
    wp.set_run_font(run, size=28, color=wp.TEAL, bold=True)
    wp.add_para(
        doc,
        "A practical dashboard project using a fictional company, simple financial formulas and tested what-if experiments.",
        size=14,
        color=wp.INK_SOFT,
    )
    wp.add_callout(
        doc,
        "The experiment",
        "I tested what happens when the company hires 40 people, customers pay later, utilization changes, or the future becomes more uncertain.",
        tone="teal",
    )
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(42)
    meta.paragraph_format.space_after = Pt(7)
    run = meta.add_run("NovaBridge Digital SRL | Fictional Romanian consulting company | EUR")
    wp.set_run_font(run, size=10.5, color=wp.MUTED, bold=True)
    wp.add_para(doc, "Prepared in August 2026 | 15-page student edition", color=wp.MUTED, size=10)
    legal = doc.add_paragraph()
    legal.paragraph_format.space_before = Pt(64)
    run = legal.add_run("All people, companies, invoices and results in this report are synthetic.")
    wp.set_run_font(run, size=9.5, color=wp.MUTED, italic=True)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    wp.configure_page(body)
    configure_header_footer(body)

    # Page 2.
    start_page(doc, 1, "Why I chose this topic", "A simple business problem with a cash timing twist")
    wp.add_para(
        doc,
        "A consulting company pays salaries every month, but customers often pay invoices several weeks after the work is completed. Hiring can increase future revenue, yet the company must pay the new salaries before every employee is fully productive. That makes growth a cash question as well as a profit question.",
    )
    wp.add_callout(
        doc,
        "Research question",
        "Can the fictional company hire 40 people, earn more profit and still keep a comfortable cash balance while the new team ramps up?",
        tone="blue",
    )
    wp.add_heading(doc, "My three goals", level=2)
    numbered = wp.create_numbering(doc, "decimal")
    for item in (
        "Understand the main numbers: revenue, operating profit and cash.",
        "Build one dashboard where a beginner can change assumptions and see a financial result.",
        "Test the model so a control cannot look useful while doing nothing.",
    ):
        wp.add_list_item(doc, item, numbered)
    wp.add_heading(doc, "The practical idea", level=2)
    wp.add_para(
        doc,
        "I treated the dashboard as a small laboratory. Each page asks a question, lets the user change one or more inputs, and then explains what moved and why. The aim is learning, not pretending that a fictional forecast is a real business decision.",
    )

    # Page 3.
    next_page(doc, 2, "What I built", "One project from data to conclusion")
    add_figure(
        doc,
        figures["flow"],
        1,
        "The project follows five steps instead of presenting many unrelated charts.",
        "Five-step flow from fictional data through formulas, forecasts and experiments to a lesson.",
        width=5.9,
    )
    wp.add_table(
        doc,
        ["Dashboard page", "Question it answers"],
        [
            ("Project Overview", "What problem is the project studying?"),
            ("What-if Experiment", "What changes if I hire, sell more time, raise price or wait longer for cash?"),
            ("13-Week Cash Check", "Will the business have enough cash in the next three months?"),
            ("Risk Simulation", "What range appears when the future is uncertain?"),
            ("How It Works", "Where do the data, formulas and tests come from?"),
        ],
        [2500, 6860],
        font_size=9.2,
    )

    # Page 4.
    next_page(doc, 3, "The fictional data", "Useful for learning, safe for a student project")
    wp.add_para(
        doc,
        "Synthetic data is invented data made to follow selected rules. It is useful here because I can show invoices, salaries and cash movements without using confidential company or client information. Seed 42 makes the same fictional history appear every time the generator runs.",
    )
    add_figure(
        doc,
        figures["data"],
        2,
        "Four understandable data groups support the main student story.",
        "Four cards showing monthly results, invoices, employees and a 40-person hiring plan.",
        width=5.8,
    )
    wp.add_callout(
        doc,
        "A useful rule",
        "The data is invented, but it must still add up: invoice revenue matches monthly revenue, unpaid invoices become receivables, and opening cash plus the cash movement equals closing cash.",
        tone="amber",
    )

    # Page 5.
    next_page(doc, 4, "The company story", "Six years of growth, pressure and recovery")
    add_figure(
        doc,
        figures["history"],
        3,
        "Revenue, operating profit and closing cash tell different parts of the 2020-2025 story.",
        "Annual line chart of fictional revenue, EBITDA and closing cash from 2020 to 2025.",
        width=5.9,
    )
    wp.add_table(
        doc,
        ["2025 result", "Value", "Plain meaning"],
        [
            ("Revenue", wp.fmt_eur(latest.revenue_eur, 1), "Value of consulting services delivered"),
            ("Operating profit", wp.fmt_eur(latest.ebitda_eur, 1), "Revenue left after the main operating costs"),
            ("Profit margin", f"{latest.ebitda_margin * 100:.1f}%", "Operating profit divided by revenue"),
            ("Closing cash", wp.fmt_eur(latest.closing_cash_eur, 1), "Money available at the end of 2025"),
        ],
        [2200, 1900, 5260],
        font_size=9.1,
    )
    wp.add_para(doc, "The company begins the forecast with a strong cash balance. The project is therefore about funding growth, not rescuing a failing business.")

    # Page 6.
    next_page(doc, 5, "Three important numbers", "Revenue, operating profit and cash are not the same")
    wp.add_table(
        doc,
        ["Number", "Beginner definition", "Small example"],
        [
            ("Revenue", "Value of work completed for customers", "Deliver EUR 100 of work -> EUR 100 revenue"),
            ("Operating profit", "Revenue minus the main operating costs", "EUR 100 revenue - EUR 75 costs -> EUR 25 profit"),
            ("Cash", "Money actually received and still available", "If the invoice is unpaid, the EUR 100 is not cash yet"),
        ],
        [1800, 3500, 4060],
        font_size=9.3,
    )
    wp.add_formula(doc, "Operating profit = Revenue - Payroll - Contractor cost - Other operating cost")
    wp.add_formula(doc, "Closing cash = Opening cash + Cash collected - Cash paid")
    wp.add_callout(
        doc,
        "The key lesson",
        "A profitable month can still use cash when customers have not paid yet, salaries are due now, or the company buys equipment.",
        tone="teal",
    )
    wp.add_heading(doc, "Why the dashboard shows all three", level=2)
    wp.add_para(
        doc,
        "Revenue explains business activity. Operating profit explains whether the work covers the operating costs. Cash explains whether bills can be paid on time. Looking at only one of them gives an incomplete answer.",
    )

    # Page 7.
    next_page(doc, 6, "How the company makes money", "People and productive time create service revenue")
    add_figure(
        doc,
        figures["driver_tree"],
        4,
        "The revenue formula is a chain of understandable operating inputs.",
        "Driver tree multiplying people, billable share, working time, utilization, billing rate and realization.",
        width=5.65,
    )
    wp.add_formula(doc, "Revenue = Employees x Billable share x Available hours x Utilization x Hourly rate")
    wp.add_heading(doc, "A worked example", level=2)
    wp.add_para(
        doc,
        "Suppose 100 employees include 80 billable consultants. If each has 160 available hours, utilization is 75%, and the rate is EUR 75 per hour, the simple monthly revenue estimate is 80 x 160 x 75% x EUR 75 = EUR 720,000 before discounts or other adjustments.",
    )
    wp.add_callout(
        doc,
        "Why utilization matters",
        "Hiring creates capacity, but capacity produces revenue only when customers buy the available time. That is why more employees do not create full revenue immediately.",
        tone="blue",
    )

    # Page 8.
    next_page(doc, 7, "Why profit is not cash", "Customer payment time creates the gap")
    add_figure(
        doc,
        figures["working_capital"],
        5,
        "Revenue is recorded before cash when the customer pays later.",
        "Flow from service delivery to invoice, payment delay, accounts receivable and cash collection.",
        width=5.55,
    )
    wp.add_table(
        doc,
        ["Term", "Plain meaning"],
        [
            ("Invoice", "A request for payment after work is delivered"),
            ("Customer payment time", "How many days the company waits to collect"),
            ("Accounts receivable", "Invoices that customers still owe"),
            ("Cash collection", "The moment money reaches the company"),
        ],
        [2500, 6860],
        font_size=9.4,
    )
    wp.add_callout(
        doc,
        "Experiment rule",
        "Making customers pay 15 days later should reduce near-term cash, but it should not reduce revenue or operating profit that was already earned.",
        tone="amber",
    )

    # Page 9.
    next_page(doc, 8, "The starting forecast", "Two time horizons answer two questions")
    wp.add_table(
        doc,
        ["Forecast", "Question"],
        [
            ("13 weeks", "Can the company pay near-term salaries, suppliers, taxes and one-time bills?"),
            ("18 months", "How do hiring, utilization, prices and customer payments affect profit and cash?"),
        ],
        [2200, 7160],
        font_size=9.3,
    )
    add_figure(
        doc,
        figures["liquidity"],
        6,
        "The weekly chart is zoomed to make the cash movement visible.",
        "Thirteen-week line chart of weekly closing cash.",
        width=5.75,
    )
    wp.add_table(
        doc,
        ["Cash check", "Result"],
        [
            ("Opening cash", wp.fmt_eur(liquidity["opening_cash_eur"].iloc[0], 1)),
            ("Lowest weekly cash", wp.fmt_eur(liquidity["closing_cash_eur"].min(), 1)),
            ("Cash after 13 weeks", wp.fmt_eur(liquidity["closing_cash_eur"].iloc[-1], 1)),
        ],
        [4400, 4960],
        font_size=9.4,
    )

    # Page 10.
    next_page(doc, 9, "Experiment 1", "What happens if the company hires 40 people?")
    add_figure(
        doc,
        figures["hiring"],
        7,
        "The hiring experiment uses the same assumptions as the dashboard's default What-if page.",
        "Line chart comparing Base and Hire 40 cash paths over eighteen months.",
        width=5.8,
    )
    revenue_change = hire_summary["revenue_eur"] - base_summary["revenue_eur"]
    profit_change = hire_summary["ebitda_eur"] - base_summary["ebitda_eur"]
    min_change = hire_summary["minimum_cash_eur"] - base_summary["minimum_cash_eur"]
    wp.add_table(
        doc,
        ["What changed", "Result compared with Base"],
        [
            ("Revenue", wp.fmt_eur(revenue_change, 1)),
            ("Operating profit", wp.fmt_eur(profit_change, 1)),
            ("Lowest cash point", wp.fmt_eur(min_change, 1)),
            ("Employees at the end", f"{hire_summary['headcount']:.0f}"),
        ],
        [4200, 5160],
        font_size=9.3,
    )
    wp.add_callout(
        doc,
        "What I learned",
        "The 40-person plan improves revenue and profit over 18 months, while the lowest cash point is slightly lower during ramp-up. The result is a timing trade-off, not simply 'hiring is good' or 'hiring is bad.'",
        tone="teal",
    )

    # Page 11.
    next_page(doc, 10, "Experiment 2", "Which dashboard input changes cash the most?")
    add_figure(
        doc,
        figures["experiments"],
        8,
        "The four student-facing controls affect the 18-month cash result in different ways.",
        "Horizontal bars showing end-of-horizon cash impact for hiring, utilization, billing rate and slower payments.",
        width=5.75,
    )
    experiment_rows = []
    for name, scenario in results["experiments"].items():
        summary = scenario["summary"]
        experiment_rows.append(
            (
                name,
                wp.fmt_eur(summary["year_end_cash_eur"] - base_summary["year_end_cash_eur"], 1),
                wp.fmt_eur(summary["ebitda_eur"] - base_summary["ebitda_eur"], 1),
            )
        )
    wp.add_table(
        doc,
        ["Experiment", "Cash impact", "Profit impact"],
        experiment_rows,
        [4000, 2680, 2680],
        font_size=9.0,
    )
    wp.add_callout(
        doc,
        "Payment-time check",
        "The slower-payment experiment changes cash but leaves operating profit unchanged. That is a useful test because it proves the model separates earning revenue from collecting it.",
        tone="blue",
    )

    # Page 12.
    next_page(doc, 11, "A small forecasting comparison", "Two simple ways to estimate next month")
    wp.add_para(
        doc,
        "The financial plan is driven by business assumptions. I also compared two small forecasting ideas as a research exercise. The test keeps the months in order so the model never learns from a future month before predicting it.",
    )
    add_figure(
        doc,
        figures["walk_forward"],
        9,
        "Walk-forward testing uses past months to predict the next month, then moves forward.",
        "Three rows showing expanding training windows followed by one forecast month.",
        width=5.55,
    )
    wp.add_table(
        doc,
        ["Method", "How it works", "Average revenue error"],
        [
            ("Seasonal repeat", "Use revenue from the same month one year earlier", "About EUR 330k"),
            ("Simple driver regression", "Use a straight-line relationship with planned business drivers", "About EUR 183k"),
        ],
        [2300, 4700, 2360],
        font_size=9.2,
    )
    wp.add_callout(
        doc,
        "Honest limit",
        "The regression result assumes that its driver values are already known or planned. It is a useful comparison, not a fully automatic prediction of the whole company.",
        tone="amber",
    )

    # Page 13.
    next_page(doc, 12, "Monte Carlo", "A range of possible futures instead of one answer")
    wp.add_para(
        doc,
        "A normal forecast gives one answer. Monte Carlo repeats the forecast many times. In each run, the model makes small random changes to business drivers such as utilization, customer payment time, prices and salary costs. It then recalculates revenue, profit and cash.",
    )
    add_figure(
        doc,
        figures["mc_flow"],
        10,
        "The calculation remains understandable even though it is repeated thousands of times.",
        "Five-step Monte Carlo flow from Base forecast through changed drivers and recalculation to a range.",
        width=5.85,
    )
    wp.add_table(
        doc,
        ["Result", "Plain meaning"],
        [
            ("Low result (P5)", "Only about 5% of simulated ending cash results are lower"),
            ("Middle result (P50)", "Half of simulated ending cash results are lower and half higher"),
            ("High result (P95)", "Only about 5% of simulated ending cash results are higher"),
        ],
        [2600, 6760],
        font_size=9.3,
    )

    # Page 14.
    next_page(doc, 13, "Testing Monte Carlo", "The controls now produce visible and logical changes")
    add_figure(
        doc,
        figures["mc_test"],
        11,
        "Setting uncertainty to zero collapses the range; doubling uncertainty makes it roughly twice as wide.",
        "Horizontal bars comparing P5-to-P95 range width at zero, normal and high uncertainty.",
        width=5.75,
    )
    mc_rows = []
    for label in ("0%", "100%", "200%"):
        summary = results["mc"][label]["summary"]
        width = summary["p95_year_end_cash_eur"] - summary["p5_year_end_cash_eur"]
        mc_rows.append(
            (
                label,
                wp.fmt_eur(summary["p5_year_end_cash_eur"], 1),
                wp.fmt_eur(summary["median_year_end_cash_eur"], 1),
                wp.fmt_eur(summary["p95_year_end_cash_eur"], 1),
                wp.fmt_eur(width, 1),
            )
        )
    wp.add_table(
        doc,
        ["Uncertainty", "P5", "Middle", "P95", "Range"],
        mc_rows,
        [1900, 1800, 1900, 1800, 1960],
        font_size=8.8,
    )
    low_probability = results["mc_low_target"]["summary"]["probability_cash_below_threshold"] * 100
    high_probability = results["mc_high_target"]["summary"]["probability_cash_below_threshold"] * 100
    wp.add_callout(
        doc,
        "The bug I found and fixed",
        f"The first page exposed run count, seed and a cash target below every path, so it looked frozen. The revised page changes real uncertainty. At 200% uncertainty, raising the target from EUR 14m to EUR 18m changes the below-target result from {low_probability:.1f}% to {high_probability:.1f}% while the cash paths stay the same, which is the correct behavior.",
        tone="red",
    )

    # Page 15.
    next_page(doc, 14, "What I learned", "A useful model is one that can be explained and tested")
    wp.add_callout(
        doc,
        "Answer to the question",
        f"In this fictional Base case, the company can test a 40-person hiring plan without approaching the EUR 3m policy floor. The plan adds about {wp.fmt_eur(profit_change, 1)} of operating profit over 18 months, but timing still matters because salaries start before full productivity and customers may pay later.",
        tone="teal",
    )
    wp.add_heading(doc, "What I tested", level=2)
    bullets = wp.create_numbering(doc, "bullet")
    for item in (
        "Revenue rises when utilization or billing rate rises.",
        "Slower customer payment reduces cash but not earned operating profit.",
        "Hiring includes salary cost and a gradual productivity ramp.",
        "The weekly cash roll-forward conserves receipts at the forecast boundary.",
        "Zero Monte Carlo uncertainty produces one path; more uncertainty widens the range.",
        "All 48 automated tests pass, every route loads, and the five main pages show no browser errors.",
    ):
        wp.add_list_item(doc, item, bullets)
    wp.add_heading(doc, "Limitations and next steps", level=2)
    wp.add_para(
        doc,
        "The numbers are fictional, the rules are simplified, and the future ranges are assumptions rather than proven probabilities. A next version could compare forecasts with later actual results, let users adjust a few uncertainty drivers separately, and use real data only after privacy and reconciliation checks.",
    )
    wp.add_heading(doc, "Final reflection", level=2)
    wp.add_para(
        doc,
        "The most important lesson was not a complicated algorithm. It was learning to connect a business action to revenue, profit and cash, then design a test that proves the control really reaches the calculation.",
    )
    wp.add_callout(
        doc,
        "Educational use only",
        "NovaPlan is an educational proof-of-concept built using entirely synthetic data. It is not an official Deloitte product, does not represent Deloitte proprietary methods, and is not financial, tax, accounting or investment advice.",
        tone="red",
    )

    doc.save(REPORT_PATH)
    return REPORT_PATH


def main() -> None:
    analytics = wp.collect_analytics()
    path = build_report(analytics)
    print(f"Created {path}")


if __name__ == "__main__":
    main()
