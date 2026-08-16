"""Build the 15-page beginner-friendly NovaPlan project summary.

The summary reuses the audited live model outputs and publication figures from
the full white-paper builder while presenting a shorter teaching narrative.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

import build_whitepaper as wp


SUMMARY_PATH = wp.OUT_DIR / "NovaPlan_15_Page_Project_Summary.docx"


def configure_summary_header_footer(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = paragraph.add_run("NOVAPLAN | PROJECT SUMMARY")
    wp.set_run_font(left, size=8.5, color=wp.MUTED, bold=True)
    right = paragraph.add_run("\t15-PAGE EDITION | AUGUST 2026")
    wp.set_run_font(right, size=8.5, color=wp.MUTED, bold=True)
    wp.set_paragraph_border(paragraph, wp.GRID, side="bottom", size="8")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("Beginner-friendly synthetic-data decision guide")
    wp.set_run_font(left, size=8.5, color=wp.MUTED)
    page = paragraph.add_run("\tPage ")
    wp.set_run_font(page, size=8.5, color=wp.MUTED)
    wp.add_page_field(paragraph)


def start_page(doc: Document, kicker: str, title: str) -> None:
    wp.add_section_kicker(doc, kicker)
    wp.add_heading(doc, title, level=1)


def next_page(doc: Document, kicker: str, title: str) -> None:
    wp.page_break(doc)
    start_page(doc, kicker, title)


def add_captioned_figure(
    doc: Document,
    path: Path,
    number: int,
    caption: str,
    alt_text: str,
    width: float = 5.85,
) -> None:
    wp.add_figure(
        doc,
        path,
        f"Figure {number}. {caption}",
        width=width,
        alt_text=alt_text,
    )


def scenario_rows(cases: pd.DataFrame) -> list[tuple[str, str, str, str, str]]:
    rows: list[tuple[str, str, str, str, str]] = []
    for row in cases.itertuples():
        rows.append(
            (
                str(row.scenario),
                wp.fmt_eur(row.revenue_eur, 2),
                wp.fmt_eur(row.ebitda_eur, 2),
                wp.fmt_eur(row.year_end_cash_eur, 2),
                wp.fmt_eur(row.minimum_cash_eur, 2),
            )
        )
    return rows


def build_summary(analytics: dict[str, object], figures: dict[str, Path]) -> Path:
    monthly: pd.DataFrame = analytics["monthly"]
    annual: pd.DataFrame = analytics["annual"]
    datasets: dict[str, pd.DataFrame] = analytics["datasets"]
    liquidity: pd.DataFrame = analytics["liquidity"]
    cases: pd.DataFrame = analytics["decision_cases"].copy()
    cases["scenario"] = cases["scenario"].replace({"2026 Expansion Decision": "Expansion"})
    sensitivity: pd.DataFrame = analytics["sensitivity"]
    monte_carlo: dict[str, object] = analytics["monte_carlo"]
    mc_summary = monte_carlo["summary"]

    latest = monthly.iloc[-1]
    ltm = annual.iloc[-1]
    base = cases.loc[cases["scenario"].eq("Base")].iloc[0]
    expansion = cases.loc[cases["scenario"].eq("Expansion")].iloc[0]
    risk = cases.loc[cases["scenario"].eq("Risk")].iloc[0]
    mitigation = cases.loc[cases["scenario"].eq("Mitigation")].iloc[0]

    doc = Document()
    wp.configure_styles(doc)
    wp.configure_page(doc.sections[0])
    doc.core_properties.title = "NovaPlan: 15-Page Project Summary"
    doc.core_properties.subject = "Beginner-friendly summary of the business-planning and liquidity simulator"
    doc.core_properties.author = "NovaPlan project team"
    doc.core_properties.keywords = "business planning, liquidity, KPI education, forecasting, Monte Carlo, synthetic data"
    doc.core_properties.comments = "Condensed edition generated from the checked-in seed-42 synthetic model outputs."

    # Page 1 - editorial cover.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(78)
    spacer.paragraph_format.space_after = Pt(0)
    wp.add_section_kicker(doc, "15-page project summary | Condensed teaching edition")
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("NovaPlan")
    wp.set_run_font(run, size=34, color=wp.TEAL, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("How the dashboard turns business drivers into profit, cash and risk decisions")
    wp.set_run_font(run, size=20, color=wp.INK, bold=True)
    wp.add_para(
        doc,
        "A concise guide to the synthetic company, core KPIs, forecasting methods, expansion scenarios and Monte Carlo simulation.",
        size=13,
        color=wp.INK_SOFT,
    )
    wp.add_callout(
        doc,
        "The decision",
        "Can a fictional consulting company hire 40 people, create value and remain comfortably liquid while demand, productivity and collection timing remain uncertain?",
        tone="teal",
    )
    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(34)
    meta.paragraph_format.space_after = Pt(8)
    run = meta.add_run("NovaBridge Digital SRL | Fictional Romanian technology consultancy | EUR reporting")
    wp.set_run_font(run, size=10.5, color=wp.MUTED, bold=True)
    wp.add_para(doc, "Prepared 12 August 2026", color=wp.MUTED, size=10)
    legal = doc.add_paragraph()
    legal.paragraph_format.space_before = Pt(58)
    run = legal.add_run("All companies, people, customers, invoices and financial observations in this summary are synthetic.")
    wp.set_run_font(run, size=9, color=wp.MUTED, italic=True)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    wp.configure_page(body)
    configure_summary_header_footer(body)

    # Page 2 - project overview.
    start_page(doc, "1 | Project overview", "One connected decision system")
    wp.add_para(
        doc,
        "NovaPlan is an educational business-planning and liquidity simulator. It does not treat charts as isolated outputs. The same driver chain connects people, productive capacity, revenue, operating costs, profit, working capital and cash. Forecasts, scenarios, sensitivity tests and Monte Carlo simulations all reuse that chain.",
    )
    add_captioned_figure(
        doc,
        figures["architecture"],
        1,
        "Operating assumptions flow through one financial model before management choices and risk are compared.",
        "Flow diagram from operating drivers through revenue, profit, working capital and cash to forecasts, scenarios and decisions.",
        width=5.95,
    )
    wp.add_callout(
        doc,
        "Central lesson",
        "Revenue, profit and cash are related but not interchangeable. Hiring can improve future capacity while using cash today; slower collections can reduce cash without changing earned revenue or EBITDA.",
        tone="blue",
    )
    wp.add_table(
        doc,
        ["Planning lens", "Question answered"],
        [
            ("Actual KPIs", "What has happened in the synthetic business?"),
            ("Forecast", "Where is the current operating plan heading?"),
            ("Scenario", "What happens if a coherent set of assumptions changes?"),
            ("Risk", "How wide is the range when uncertain drivers vary together?"),
        ],
        [2200, 7160],
        font_size=8.8,
    )

    # Page 3 - synthetic data.
    next_page(doc, "2 | Synthetic evidence", "A fictional company with internally connected records")
    wp.add_para(
        doc,
        "The data represents NovaBridge Digital SRL from January 2020 through December 2025. Seed 42 makes the generated history repeatable. The generator creates a realistic teaching story - pandemic pressure, recovery, scale-up, slower demand and renewed growth - without using confidential information.",
    )
    add_captioned_figure(
        doc,
        figures["ecosystem"],
        2,
        "Eight linked datasets provide the evidence behind the monthly financial backbone.",
        "Data ecosystem showing monthly KPIs linked to invoices, customers, employees, hiring, operating costs, capital expenditure and weekly cash.",
        width=5.35,
    )
    wp.add_table(
        doc,
        ["Evidence group", "Records", "Purpose"],
        [
            ("Monthly financial backbone", f"{len(datasets['monthly_kpis'])} months x {datasets['monthly_kpis'].shape[1]} fields", "Drivers, profit, working capital and cash"),
            ("Customer evidence", f"{len(datasets['invoices']):,} invoices; {len(datasets['customers'])} customers", "Revenue, terms, payment dates and concentration"),
            ("Workforce evidence", f"{len(datasets['employees'])} active employees; 40-hire plan", "Headcount, salary, role mix and joining cohorts"),
            ("Costs and liquidity", f"{len(datasets['opex'])} OPEX rows; {len(datasets['capex'])} assets; {len(datasets['weekly_cash_actuals'])} weeks", "Expense, investment and cash timing"),
        ],
        [2400, 2600, 4360],
        font_size=8.3,
    )
    wp.add_callout(
        doc,
        "What synthetic means",
        "The numbers are invented but economically connected: invoice totals reconcile to monthly revenue, receivables reconcile to collections, and cash rolls from opening to closing balances.",
        tone="amber",
    )
    wp.add_table(
        doc,
        ["Parameter family", "Illustrative settings", "Scope"],
        [
            ("Starting position", "272 people; EUR 5.20m backlog; EUR 4.70m cash", "January 2020 synthetic opening state"),
            ("Workforce", "New-hire ramp 35% / 65% / 85% / 100%; 2.25% employer on-cost", "Historical generator only"),
            ("Commercial", "96.5% realization; 1.42x contractor markup; 30 / 45 / 60-day terms", "Synthetic revenue and invoice logic"),
            ("Planning", "18 months; 13 weeks; 10,000 runs; seed 42; EUR 3m floor", "Dashboard defaults, not universal truths"),
        ],
        [1900, 4300, 3160],
        font_size=7.7,
    )

    # Page 4 - historical story.
    next_page(doc, "3 | Synthetic company story", "Six years of growth, stress and recovery")
    add_captioned_figure(
        doc,
        figures["history"],
        3,
        "Revenue, EBITDA and closing cash describe different dimensions of the same six-year history.",
        "Line chart of annual synthetic revenue, EBITDA and closing cash from 2020 through 2025.",
        width=5.95,
    )
    wp.add_table(
        doc,
        ["2025 headline", "Result", "Beginner reading"],
        [
            ("Revenue", wp.fmt_eur(ltm.revenue_eur, 2), "Value of services earned"),
            ("EBITDA", wp.fmt_eur(ltm.ebitda_eur, 2), f"{ltm.ebitda_margin * 100:.1f}% operating margin"),
            ("Operating cash flow", wp.fmt_eur(ltm.operating_cash_flow_eur, 2), "Cash generated after working-capital timing"),
            ("Closing cash", wp.fmt_eur(ltm.closing_cash_eur, 2), "Liquidity at December 2025"),
            ("Headcount", f"{ltm.headcount:.0f}", "People employed at the cut-off"),
        ],
        [2200, 1900, 5260],
        font_size=8.6,
    )
    wp.add_para(
        doc,
        f"The latest year records {wp.fmt_eur(ltm.revenue_eur, 2)} of revenue and ends with {wp.fmt_eur(ltm.closing_cash_eur, 2)} of cash. This is a comfortable starting position, so the expansion case is a growth decision rather than a rescue plan.",
    )
    wp.add_callout(
        doc,
        "Reading caution",
        "Synthetic history can demonstrate mechanics and tests, but it cannot prove real-world forecasting performance or causal relationships.",
        tone="red",
    )

    # Page 5 - revenue drivers.
    next_page(doc, "4 | Operating KPIs", "How people and productive time become revenue")
    add_captioned_figure(
        doc,
        figures["driver_tree"],
        4,
        "A transparent driver tree explains where service revenue comes from.",
        "Driver tree multiplying headcount, billable share, working days, hours, utilization, billing rate and realization to obtain revenue.",
        width=5.75,
    )
    wp.add_formula(doc, "Capacity hours = Headcount x Billable share x Working days x Hours/day")
    wp.add_formula(doc, "Revenue = Capacity hours x Utilization x Billing rate x Realization + backlog and contractor revenue")
    wp.add_table(
        doc,
        ["KPI", "Plain meaning", "Why it matters"],
        [
            ("Headcount", "People employed", "Creates payroll cost and future capacity"),
            ("Billable share", "People whose work can be sold", "Separates delivery roles from support roles"),
            ("Utilization", "Share of billable capacity used", "Central productivity and volume lever"),
            ("Billing rate", "Nominal price per hour", "Commercial pricing lever"),
            ("Realization", "Value retained after discounts", "Separates nominal from recognized value"),
            ("Backlog", "Contracted work not yet recognized", "Provides future revenue cover"),
        ],
        [1700, 3300, 4360],
        font_size=8.3,
    )
    wp.add_callout(
        doc,
        "Trade-off",
        "Higher utilization usually improves profit, but extremely high utilization may be difficult to sustain and can reduce capacity for training, sales or innovation.",
        tone="amber",
    )

    # Page 6 - profit to cash.
    next_page(doc, "5 | Profit and liquidity KPIs", "Why profit and cash can move differently")
    add_captioned_figure(
        doc,
        figures["working_capital"],
        5,
        "Accounts receivable carries the timing gap between service delivery and customer payment.",
        "Sequence from service delivery to invoice, payment delay and cash collection, explaining the role of DSO and accounts receivable.",
        width=5.55,
    )
    wp.add_formula(doc, "EBITDA = Revenue - Payroll - Contractor cost - Other OPEX")
    wp.add_formula(doc, "Operating cash flow = EBITDA - Tax - Change in AR + Change in AP")
    wp.add_formula(doc, "Free cash flow = Operating cash flow - Capex")
    wp.add_table(
        doc,
        ["KPI", "Plain meaning", "Management use"],
        [
            ("EBITDA", "Core operating profit before depreciation and financing/tax", "Compares operating profitability"),
            ("AR", "Customer amounts still unpaid", "Shows cash tied up after revenue is earned"),
            ("DSO", "Approximate collection time", "Higher values normally reduce near-term cash"),
            ("AP / DPO", "Supplier amounts and payment time", "Longer timing preserves cash temporarily"),
            ("Operating cash flow", "Cash generated by operations", "Tests whether profit converts to cash"),
            ("Free cash flow", "Cash after capital investment", "Shows cash available before financing/distributions"),
        ],
        [1900, 3900, 3560],
        font_size=8.1,
    )
    wp.add_callout(
        doc,
        "Capex lesson",
        "Capital expenditure uses cash immediately but does not reduce EBITDA. Depreciation reduces EBIT later without repeating the original cash payment.",
        tone="blue",
    )

    # Page 7 - forecast horizons.
    next_page(doc, "6 | Forecast horizons", "Near-term liquidity and longer-term planning answer different questions")
    wp.add_table(
        doc,
        ["Forecast", "Grain", "Primary question"],
        [
            ("13-week direct cash", "Week", "Can the company meet near-term obligations and stay above its cash floor?"),
            ("18-month driver forecast", "Month", "What do demand, hiring, price and collections do to profit and cash?"),
        ],
        [2400, 1100, 5860],
        font_size=8.6,
    )
    add_captioned_figure(
        doc,
        figures["liquidity"],
        6,
        "The Base 13-week cash forecast remains well above the EUR 3m policy floor.",
        "Line chart of weekly closing cash over thirteen weeks with a horizontal EUR 3 million policy floor.",
        width=5.75,
    )
    wp.add_table(
        doc,
        ["Liquidity measure", "Base result"],
        [
            ("Opening cash", wp.fmt_eur(liquidity["opening_cash_eur"].iloc[0], 2)),
            ("13-week collections", wp.fmt_eur(liquidity["customer_collections_eur"].sum(), 2)),
            ("13-week disbursements", wp.fmt_eur(liquidity["total_disbursements_eur"].sum(), 2)),
            ("Lowest weekly cash", wp.fmt_eur(liquidity["closing_cash_eur"].min(), 2)),
            ("Forecast-end cash", wp.fmt_eur(liquidity["closing_cash_eur"].iloc[-1], 2)),
        ],
        [4300, 5060],
        font_size=8.7,
    )
    wp.add_callout(
        doc,
        "Why two horizons",
        "Weekly detail is better for payment timing. Monthly detail is better for staffing, pricing, demand and profitability. Neither forecast replaces the other.",
        tone="teal",
    )

    # Page 8 - scenario decision.
    next_page(doc, "7 | Scenario decision", "Expansion creates profit but temporarily uses liquidity")
    add_captioned_figure(
        doc,
        figures["scenario_paths"],
        7,
        "Four coherent management stories remain above the EUR 3m floor but separate materially over time.",
        "Line chart comparing monthly cash paths for Base, Expansion, Risk and Mitigation cases through June 2027.",
        width=5.75,
    )
    wp.add_table(
        doc,
        ["Case", "Revenue", "EBITDA", "Jun 2027 cash", "Minimum cash"],
        scenario_rows(cases),
        [1450, 1850, 1850, 2100, 2110],
        font_size=8.0,
    )
    wp.add_callout(
        doc,
        "Decision reading",
        f"Expansion adds {wp.fmt_eur(expansion.revenue_eur - base.revenue_eur, 2)} of revenue and {wp.fmt_eur(expansion.ebitda_eur - base.ebitda_eur, 2)} of EBITDA, but minimum cash falls by {wp.fmt_eur(base.minimum_cash_eur - expansion.minimum_cash_eur, 2)} during ramp-up. Mitigation produces the strongest combined profit-and-cash result.",
        tone="teal",
    )
    wp.add_para(
        doc,
        "A scenario is deterministic: each case answers an explicit what-if question. It is not a probability-weighted prediction.",
    )

    # Page 9 - sensitivity.
    next_page(doc, "8 | Sensitivity analysis", "One lever at a time reveals financial materiality")
    add_captioned_figure(
        doc,
        figures["sensitivity"],
        8,
        "Utilization and billing rate have the largest forecast-end cash effects among the tested shocks.",
        "Horizontal sensitivity chart showing positive and negative forecast-end cash effects for utilization, billing rate, salary, DSO, foreign exchange, attrition and capex.",
        width=5.8,
    )
    top = sensitivity.loc[sensitivity["direction"].eq("High")].sort_values("rank")
    selected = top.loc[top["driver"].isin(["Utilization", "Billing Rate", "Salary Cost", "DSO", "Capex"])]
    rows = []
    for row in selected.itertuples():
        rows.append((row.driver, row.shock, wp.fmt_eur(row.year_end_cash_impact_eur, 2), wp.fmt_eur(row.ebitda_impact_eur, 2)))
    wp.add_table(
        doc,
        ["Driver", "High shock", "Cash impact", "EBITDA impact"],
        rows,
        [2200, 1800, 2680, 2680],
        font_size=8.4,
    )
    wp.add_callout(
        doc,
        "DSO lesson",
        "Ten days slower collections reduces forecast-end cash by about EUR 1.43m while EBITDA is unchanged. Collection timing changes cash and receivables, not earned revenue.",
        tone="blue",
    )
    wp.add_table(
        doc,
        ["Method", "Question"],
        [
            ("Scenario", "What if several named assumptions occur together?"),
            ("Sensitivity", "Which tested lever has the largest effect?"),
            ("Monte Carlo", "What distribution appears when uncertain drivers vary together?"),
        ],
        [2200, 7160],
        font_size=8.5,
    )

    # Page 10 - algorithms.
    next_page(doc, "9 | Forecasting algorithms", "Complexity must earn its place")
    wp.add_para(
        doc,
        "The statistical diagnostics benchmark one-month-ahead estimates. They do not automatically replace the driver-based financial plan. Their governance purpose is to test whether a more complex model improves out-of-sample accuracy enough to justify harder explanation and maintenance.",
    )
    add_captioned_figure(
        doc,
        figures["walk_forward"],
        9,
        "Walk-forward validation trains only on information available before each predicted month.",
        "Expanding-window time-series validation diagram with successive training blocks and one next-month prediction in each fold.",
        width=5.55,
    )
    wp.add_table(
        doc,
        ["Algorithm", "How it works", "Main trade-off"],
        [
            ("Seasonal Naive", "Repeats the same month one year earlier", "Transparent but ignores current drivers"),
            ("Driver Regression", "Ordinary multivariate linear regression", "Interpretable but assumes stable linear effects"),
            ("ElasticNet", "Regularized and standardized linear model", "Stable with correlated drivers but still linear"),
            ("Gradient Boosting", "Conservative nonlinear decision trees", "Learns interactions but can overfit 72 months"),
        ],
        [1900, 4100, 3360],
        font_size=8.2,
    )
    wp.add_formula(doc, "WAPE = sum(abs(Actual - Predicted)) / sum(abs(Actual))", "Lower WAPE is better.")
    wp.add_callout(
        doc,
        "No leakage",
        "With 72 monthly observations and a 36-month initial training window, every model produces 36 chronological test predictions. The data is never shuffled.",
        tone="blue",
    )

    # Page 11 - algorithm results.
    next_page(doc, "10 | Model diagnostics", "ElasticNet is strongest overall on this synthetic history")
    add_captioned_figure(
        doc,
        figures["model_wape"],
        10,
        "No model wins every target, although ElasticNet has the strongest general WAPE profile.",
        "Grouped bar chart comparing WAPE for four forecasting algorithms across revenue, utilization, collections, EBITDA and closing cash.",
        width=5.85,
    )
    wp.add_table(
        doc,
        ["Revenue model", "MAE", "RMSE", "WAPE", "Direction"],
        [
            ("ElasticNet", "EUR 58.4k", "EUR 69.8k", "2.09%", "94.4%"),
            ("Gradient Boosting", "EUR 172.4k", "EUR 215.9k", "6.17%", "80.6%"),
            ("Driver Regression", "EUR 183.0k", "EUR 221.9k", "6.55%", "80.6%"),
            ("Seasonal Naive", "EUR 329.8k", "EUR 377.2k", "11.80%", "69.4%"),
        ],
        [2300, 1750, 1750, 1750, 1810],
        font_size=8.3,
    )
    wp.add_callout(
        doc,
        "Interpretation boundary",
        "Feature-based models use observed same-month driver values. The test therefore asks how well a target can be estimated when those drivers are known or planned; it is not a fully autonomous forecast of both drivers and outcomes.",
        tone="amber",
    )
    wp.add_para(
        doc,
        "MAE is the typical miss in the metric's own unit. RMSE penalizes occasional large misses. Directional accuracy tests whether rises and falls are anticipated.",
    )

    # Page 12 - Monte Carlo method.
    next_page(doc, "11 | Monte Carlo method", "The simulator samples business drivers, not cash")
    wp.add_para(
        doc,
        "A deterministic forecast gives one path. Monte Carlo repeats the model 10,000 times while uncertain drivers move together. Each run perturbs operations, rebuilds profit and working capital, and then derives one internally coherent cash path.",
    )
    add_captioned_figure(
        doc,
        figures["monte_carlo_flow"],
        11,
        "The simulation preserves an auditable path from operating uncertainty to liquidity.",
        "Five-step Monte Carlo flow from Base forecast through driver shocks and recalculation to cash paths and risk summaries.",
        width=5.75,
    )
    wp.add_table(
        doc,
        ["Driver group", "Sampling idea", "Hard boundary"],
        [
            ("Utilization", "Macro path plus monthly noise", "55%-90%"),
            ("DSO", "Faster in favorable conditions plus monthly noise", "30-95 days"),
            ("Commercial", "Billing realization, backlog conversion and demand", "Realization 90%-108%; backlog 82%-118%"),
            ("People cost", "Salary inflation and attrition shocks by run", "Salary -5 pp to +8 pp; attrition -6 pp to +8 pp"),
            ("EUR/RON", "Constant foreign-exchange shock by run", "-5% to +5%"),
        ],
        [2000, 4200, 3160],
        font_size=8.1,
    )
    wp.add_callout(
        doc,
        "Reproducibility",
        "The published run uses 10,000 simulations, seed 42 and NumPy's PCG64 bit generator. The same code, data, assumptions and software environment reproduce the same draws.",
        tone="blue",
    )

    # Page 13 - Monte Carlo result.
    next_page(doc, "12 | Monte Carlo results", "The Base forecast has a wide but comfortable liquidity range")
    add_captioned_figure(
        doc,
        figures["monte_carlo_fan"],
        12,
        "The dark band contains the middle 50% of monthly outcomes and the light band P5-P95.",
        "Monte Carlo fan chart of monthly cash with P5 to P95 and P25 to P75 bands, median line and EUR 3 million floor.",
        width=5.8,
    )
    ranking = monte_carlo["risk_driver_ranking"].head(4)
    driver_text = ", ".join(
        f"{row.driver} {row.correlation_with_year_end_cash:+.2f}" for row in ranking.itertuples()
    )
    wp.add_table(
        doc,
        ["Risk measure", "Result", "Beginner interpretation"],
        [
            ("P5 forecast-end cash", wp.fmt_eur(mc_summary["p5_year_end_cash_eur"], 2), "Only 5% of simulated June 2027 balances are lower"),
            ("Median / P50", wp.fmt_eur(mc_summary["median_year_end_cash_eur"], 2), "Half of outcomes are lower and half higher"),
            ("P95 forecast-end cash", wp.fmt_eur(mc_summary["p95_year_end_cash_eur"], 2), "Only 5% of balances are higher"),
            ("EUR 3m floor breach", f"{mc_summary['probability_cash_below_threshold'] * 100:.2f}%", "No breach occurred in the selected 10,000-run model"),
            ("Strongest associations", driver_text, "Correlation describes co-movement, not isolated causation"),
        ],
        [2500, 2600, 4260],
        font_size=8.0,
    )
    wp.add_callout(
        doc,
        "Zero events is not zero risk",
        "No run breaches EUR 3m under the selected distributions and synthetic opening position. This does not prove that a breach is impossible or cover risks excluded from the model.",
        tone="red",
    )
    wp.add_para(
        doc,
        "Each monthly percentile is calculated across all simulations at that month. The P5 line is not one single unlucky company path.",
    )

    # Page 14 - dashboard use and governance.
    next_page(doc, "13 | Dashboard use", "Each page answers a management question")
    wp.add_table(
        doc,
        ["Workspace", "Pages", "Question"],
        [
            ("Understand", "Overview; Learning Guide", "What is happening, and what does each metric mean?"),
            ("Plan", "Rolling Forecast; Workforce; Liquidity", "Where are we heading, and can we fund operations?"),
            ("Compare", "Scenario Lab; Sensitivity", "Which coherent story or individual lever matters most?"),
            ("Test uncertainty", "Diagnostics; Monte Carlo", "How accurate are models, and how wide is the risk range?"),
            ("Govern", "Data & Assumptions", "Where did inputs come from, and what judgment was entered?"),
        ],
        [1700, 3300, 4360],
        font_size=8.4,
    )
    wp.add_heading(doc, "Four kinds of information", level=2)
    wp.add_table(
        doc,
        ["Information type", "Meaning", "Example"],
        [
            ("Observed synthetic actual", "Generated historical record", "December 2025 revenue"),
            ("Calculated value", "Deterministic formula applied to inputs", "EBITDA, DSO or cash roll-forward"),
            ("Management assumption", "Human-selected planning input", "Rate increase, DSO change or hiring pace"),
            ("Model estimate", "Forecast or simulated output", "June 2027 cash or Monte Carlo P5"),
        ],
        [2600, 3600, 3160],
        font_size=8.5,
    )
    wp.add_heading(doc, "Governance boundaries", level=2)
    bullets = wp.create_numbering(doc, "bullet")
    for item in (
        "Runtime scenario presets are defined in Python; a separate scenarios YAML file currently differs and is not loaded by the app.",
        "Historical generator assumptions and interactive planning defaults have different scopes and must not be blended.",
        "The audit CSV records entries and rationale but does not provide authentication, approvals, immutable storage or scenario replay.",
        "The full three-statement engine and lighter interactive planner are separate MVP implementations.",
        "Default Monte Carlo wraps the Base case; scenario-specific simulation requires an explicit programmatic override.",
    ):
        wp.add_list_item(doc, item, bullets)
    wp.add_callout(
        doc,
        "Governance principle",
        "A result is useful only when the reader can identify its input data, formula, assumption owner, adjustment rationale, version and validation status.",
        tone="teal",
    )

    # Page 15 - conclusion and disclaimer.
    next_page(doc, "14 | Management conclusion", "What the project teaches and what comes next")
    wp.add_callout(
        doc,
        "Decision summary",
        f"Base planning cash reaches {wp.fmt_eur(base.year_end_cash_eur, 2)} at June 2027. Expansion is profitable but lowers the path minimum during onboarding. Risk reduces forecast-end cash to {wp.fmt_eur(risk.year_end_cash_eur, 2)}, while Mitigation improves it to {wp.fmt_eur(mitigation.year_end_cash_eur, 2)} by pacing hiring, protecting collections and adding modest pricing.",
        tone="teal",
    )
    wp.add_heading(doc, "Five questions for every decision", level=2)
    numbered = wp.create_numbering(doc, "decimal")
    for item in (
        "Which operating driver is changing: people, utilization, price, demand, collections or cost?",
        "How does that driver affect revenue and EBITDA?",
        "When does the accounting result become cash?",
        "What happens in a coherent downside or mitigation scenario?",
        "How wide is the result under explicit uncertainty, and which limitations remain outside the model?",
    ):
        wp.add_list_item(doc, item, numbered)
    wp.add_heading(doc, "Practical next steps", level=2)
    bullets = wp.create_numbering(doc, "bullet")
    for item in (
        "Unify assumptions and scenarios into one typed, versioned source of truth.",
        "Connect governed ERP, time-sheet, CRM, payroll and bank data with reconciliation controls.",
        "Calibrate forecast errors and Monte Carlo distributions from real residuals and stress periods.",
        "Add authenticated users, approvals, immutable event history and reproducible scenario snapshots.",
        "Extend the model to customer/project concentration, debt facilities, covenants and scenario-specific simulation.",
    ):
        wp.add_list_item(doc, item, bullets)
    wp.add_heading(doc, "Final perspective", level=2)
    wp.add_para(
        doc,
        "NovaPlan's strongest idea is continuity: operating choices create capacity; capacity creates revenue; profit is converted into cash through working capital; and uncertainty is applied to the same drivers before liquidity is summarized. That makes the dashboard teachable for beginners and auditable enough for serious discussion.",
    )
    wp.add_callout(
        doc,
        "Educational use only",
        "NovaPlan is an educational proof-of-concept built using entirely synthetic data. It is not an official Deloitte product, does not represent Deloitte proprietary methodologies, and should not be used as financial, tax, accounting, or investment advice.",
        tone="red",
    )
    wp.add_para(
        doc,
        "Published basis: data/monthly_kpis.csv through December 2025; 18-month forecast from January 2026 to June 2027; 10,000 Monte Carlo simulations; seed 42; EUR 3m cash floor. For full formulas, assumption registers and implementation details, consult the complete NovaPlan illustrated white paper.",
        color=wp.MUTED,
        size=9,
        italic=True,
    )

    doc.save(SUMMARY_PATH)
    return SUMMARY_PATH


def main() -> None:
    analytics = wp.collect_analytics()
    figures = wp.build_figures(analytics)
    path = build_summary(analytics, figures)
    print(f"Created {path}")


if __name__ == "__main__":
    main()
